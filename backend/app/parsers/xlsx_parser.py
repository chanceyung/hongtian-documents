"""XLSX 解析 — openpyxl"""
import hashlib
from pathlib import Path

import openpyxl

from app.models.unified_document import (
    UnifiedDocument,
    TextElement,
    ImageElement,
    TableElement,
)


class XlsxParser:

    async def parse(self, path: Path, session_id: str) -> UnifiedDocument:
        assets_dir = path.parent / "assets" / session_id
        assets_dir.mkdir(parents=True, exist_ok=True)

        workbook = openpyxl.load_workbook(str(path), data_only=True)
        warnings: list[str] = []
        texts: list[TextElement] = []
        images: list[ImageElement] = []
        tables: list[TableElement] = []

        for sheet in workbook.worksheets:
            sheet_name = sheet.title
            try:
                sheet_texts, sheet_tables, sheet_images = self._parse_sheet(
                    sheet, sheet_name, assets_dir
                )
                texts.extend(sheet_texts)
                tables.extend(sheet_tables)
                images.extend(sheet_images)
            except Exception as e:
                warnings.append(f"Sheet '{sheet_name}': {e}")

        workbook.close()

        return UnifiedDocument(
            source_file=str(path),
            source_format="xlsx",
            parse_method="openpyxl",
            total_pages=len(workbook.worksheets),
            texts=texts,
            images=images,
            tables=tables,
            parse_warnings=warnings,
        )

    def _parse_sheet(
        self, sheet, sheet_name: str, assets_dir: Path,
    ) -> tuple[list[TextElement], list[TableElement], list[ImageElement]]:
        texts: list[TextElement] = []
        tables: list[TableElement] = []
        images: list[ImageElement] = []

        data: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            row_data = [str(cell) if cell is not None else "" for cell in row]
            data.append(row_data)

        filtered = self._filter_empty(data)
        if filtered:
            headers = filtered[0]
            table_data = filtered[1:] if len(filtered) > 1 else []

            tables.append(TableElement(
                id=f"xlsx_{sheet_name}",
                page=0,
                data=table_data,
                headers=headers,
            ))

            for idx, header in enumerate(headers):
                if header.strip():
                    texts.append(TextElement(
                        id=f"xlsx_{sheet_name}_h{idx}",
                        content=header,
                        page=0,
                        level=1,
                        fingerprint=hashlib.md5(header.encode()).hexdigest(),
                    ))

        if hasattr(sheet, "_images"):
            for img_idx, img in enumerate(sheet._images):
                try:
                    img_bytes = img._data()
                    ext = getattr(img, "format", "png") or "png"
                    if ext == "jpeg":
                        ext = "jpg"
                    img_hash = hashlib.md5(img_bytes).hexdigest()[:12]
                    img_path = assets_dir / f"xlsx_{sheet_name}_img{img_idx}_{img_hash}.{ext}"
                    img_path.write_bytes(img_bytes)

                    images.append(ImageElement(
                        id=f"xlsx_{sheet_name}_img{img_idx}",
                        local_path=str(img_path),
                        page=0,
                        hash=img_hash,
                    ))
                except Exception:
                    pass

        return texts, tables, images

    def _filter_empty(self, data: list[list[str]]) -> list[list[str]]:
        if not data:
            return []

        non_empty_rows = [row for row in data if any(cell.strip() for cell in row)]
        if not non_empty_rows:
            return []

        max_cols = max(len(row) for row in non_empty_rows)
        active_cols = [
            col for col in range(max_cols)
            if any(len(row) > col and row[col].strip() for row in non_empty_rows)
        ]

        return [
            [row[col] if col < len(row) else "" for col in active_cols]
            for row in non_empty_rows
        ]
